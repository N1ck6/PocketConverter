"""
Regression tests for bugs fixed in 2.0.0 — each test names the bug it guards.
Run with: python -m pytest Test -v
"""

import helpers  # noqa: F401  (must be first: isolates %LOCALAPPDATA%)
import json
import os
import shutil
import subprocess
import sys
import tempfile
import time
import unittest
import xml.etree.ElementTree as ET
from pathlib import Path

from PIL import Image

from converter import FileConverter, main as converter_main
from converter_app import jobqueue
from converter_app.document_converter import DocumentConverter
from converter_app.formats import CONVERSION_MAP, FOLDER_TARGETS
from converter_app.utils import get_ffmpeg_exe, read_text
from helpers import docx_text, pdf_text


class TempDirTestCase(unittest.TestCase):
    def setUp(self):
        self.dir = Path(tempfile.mkdtemp(prefix='pc-test-'))

    def tearDown(self):
        shutil.rmtree(self.dir, ignore_errors=True)


class TestOutputNaming(TempDirTestCase):
    @classmethod
    def setUpClass(cls):
        cls.fc = FileConverter()

    def test_dotted_name_does_not_overwrite_other_file(self):
        src = self.dir / 'my.notes.txt'
        src.write_text('hello', encoding='utf-8')
        victim = self.dir / 'my.pdf'
        victim.write_bytes(b'KEEP ME')
        for mode in ('pdf', 'docx', 'md', 'html'):
            ok, msg = self.fc.convert_file(str(src), mode)
            self.assertTrue(ok, msg)
            self.assertTrue((self.dir / f'my.notes.{mode}').exists(), mode)
        self.assertEqual(victim.read_bytes(), b'KEEP ME')

    def test_existing_output_gets_numbered_name(self):
        src = self.dir / 'a.png'
        Image.new('RGB', (4, 4), 'red').save(src)
        (self.dir / 'a.jpg').write_bytes(b'existing')
        ok, _ = self.fc.convert_file(str(src), 'jpg')
        self.assertTrue(ok)
        self.assertEqual((self.dir / 'a.jpg').read_bytes(), b'existing')
        self.assertTrue((self.dir / 'a(1).jpg').exists())

    def test_parallel_jobs_with_same_stem_get_distinct_outputs(self):
        Image.new('RGB', (40, 40), 'red').save(self.dir / 'photo.png')
        Image.new('RGB', (40, 40), 'blue').save(self.dir / 'photo.jpg')
        results = self.fc.convert_multiple_files([str(self.dir / 'photo.png'), str(self.dir / 'photo.jpg')], 'webp')
        self.assertTrue(all(r.success for r in results), results)
        self.assertEqual(len({r.output for r in results}), 2)

    def test_unsupported_pair_is_rejected(self):
        src = self.dir / 'a.pdf'
        src.write_bytes(b'%PDF-1.4')
        ok, msg = self.fc.convert_file(str(src), 'png')
        self.assertFalse(ok)
        self.assertIn("isn't supported", msg)


class TestTextEncodings(TempDirTestCase):
    def test_cp1251_text_is_decoded_not_replaced_by_path(self):
        src = self.dir / 'ru.txt'
        src.write_bytes('Привет, мир'.encode('cp1251'))
        FileConverter().convert_file(str(src), 'md')
        self.assertEqual((self.dir / 'ru.md').read_text(encoding='utf-8').strip(), 'Привет, мир')

    def test_utf16_and_bom(self):
        (self.dir / 'u16.txt').write_bytes('текст'.encode('utf-16'))
        (self.dir / 'bom.txt').write_bytes('\ufefftext'.encode('utf-8'))
        self.assertEqual(read_text(self.dir / 'u16.txt'), 'текст')
        self.assertEqual(read_text(self.dir / 'bom.txt'), 'text')

    def test_crlf_is_not_doubled(self):
        src = self.dir / 'crlf.txt'
        src.write_bytes(b'a\r\nb\r\n')
        FileConverter().convert_file(str(src), 'md')
        self.assertEqual((self.dir / 'crlf.md').read_bytes().count(b'\r\r'), 0)


class TestDocuments(TempDirTestCase):
    def test_plain_conversion_keeps_text_unchanged(self):
        text = 'my_var * 2 * x = **not bold** and π'
        src = self.dir / 'code.txt'
        src.write_text(text, encoding='utf-8')
        fc = FileConverter()
        for mode in ('md', 'docx', 'pdf', 'html'):
            self.assertTrue(fc.convert_file(str(src), mode)[0], mode)
        self.assertEqual((self.dir / 'code.md').read_text(encoding='utf-8'), text)
        self.assertIn(text, docx_text(self.dir / 'code.docx'))
        self.assertIn('my_var * 2 * x', pdf_text(self.dir / 'code.pdf'))

    def test_txt_to_html_is_escaped(self):
        src = self.dir / 'x.txt'
        src.write_text('<script>alert(1)</script> & co', encoding='utf-8')
        FileConverter().convert_file(str(src), 'html')
        html = (self.dir / 'x.html').read_text(encoding='utf-8')
        self.assertNotIn('<script>', html)
        self.assertIn('&lt;script&gt;', html)

    def test_long_lines_wrap_inside_pdf(self):
        src = self.dir / 'long.txt'
        src.write_text('W' * 400 + '\n' + 'word ' * 200, encoding='utf-8')
        self.assertTrue(FileConverter().convert_file(str(src), 'pdf')[0])
        self.assertIn('WWWW', pdf_text(self.dir / 'long.pdf'))

    def test_scanned_pdf_gives_clear_error(self):
        import pymupdf
        pdf = pymupdf.open()
        pdf.new_page()
        pdf.save(str(self.dir / 'scan.pdf'))
        ok, msg = FileConverter().convert_file(str(self.dir / 'scan.pdf'), 'txt')
        self.assertFalse(ok)
        self.assertIn('no text layer', msg)

    def test_docx_tables_are_kept(self):
        from docx import Document
        doc = Document()
        doc.add_heading('Title', 1)
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = 'Name', 'Age'
        table.cell(1, 0).text, table.cell(1, 1).text = 'Bob', '42'
        doc.save(str(self.dir / 't.docx'))
        fc = FileConverter()
        fc.convert_file(str(self.dir / 't.docx'), 'md')
        md = (self.dir / 't.md').read_text(encoding='utf-8')
        self.assertIn('# Title', md)
        self.assertIn('| Bob | 42 |', md)

    def test_html_to_txt_drops_scripts_keeps_inline_text(self):
        (self.dir / 'p.html').write_text(
            '<html><head><script>evil()</script></head><body><p>A <b>bold</b> word</p></body></html>', encoding='utf-8')
        FileConverter().convert_file(str(self.dir / 'p.html'), 'txt')
        txt = (self.dir / 'p.txt').read_text(encoding='utf-8')
        self.assertIn('A bold word', txt)
        self.assertNotIn('evil', txt)

    def test_md_to_docx_headings_and_bullets(self):
        (self.dir / 'n.md').write_text('# Head\n\n- **one**\n- two\n\n```\ncode\n```\n', encoding='utf-8')
        self.assertTrue(FileConverter().convert_file(str(self.dir / 'n.md'), 'docx')[0])
        text = docx_text(self.dir / 'n.docx')
        self.assertIn('Head', text)
        self.assertIn('• one', text)
        self.assertNotIn('```', text)


class TestGptCleaner(unittest.TestCase):
    d = DocumentConverter()

    def test_identifiers_and_math_survive(self):
        self.assertEqual(self.d.clean_gpt_text('call my_var_name now'), 'call my_var_name now')
        self.assertEqual(self.d.clean_gpt_text('2*3*4 and 5 * 3 * 2'), '2*3*4 and 5 * 3 * 2')

    def test_emphasis_still_removed(self):
        self.assertEqual(self.d.clean_gpt_text('a *b* _c_ __d__ ~~e~~'), 'a b c d e')

    def test_code_fences_and_rules_removed(self):
        self.assertEqual(self.d.clean_gpt_text('```python\nprint(1)\n```\n---\nend'), 'print(1)\n\nend')

    def test_superscripts_and_roots(self):
        self.assertEqual(self.d._convert_formulas('x² + y³'), 'x^2 + y^3')
        self.assertEqual(self.d._convert_formulas('10⁻³'), '10^(-3)')
        self.assertEqual(self.d._convert_formulas('⁴√(16)'), '(16)^(1/4)')
        self.assertEqual(self.d._convert_formulas('√x'), '(x)^(1/2)')
        self.assertEqual(self.d._convert_formulas('x^{2n}'), 'x^(2n)')

    def test_latex_operators(self):
        self.assertEqual(self.d._convert_formulas(r'\(a \cdot b\)'), 'a * b')
        self.assertEqual(self.d._convert_formulas(r'\dfrac{1}{2} \le x'), r'(1)/(2) \le x')
        self.assertEqual(self.d._convert_formulas('a ≤ b ≠ c'), 'a <= b != c')


class TestImages(TempDirTestCase):
    @classmethod
    def setUpClass(cls):
        cls.fc = FileConverter()

    def test_transparent_png_to_jpg_has_white_background(self):
        Image.new('RGBA', (10, 10), (255, 0, 0, 0)).save(self.dir / 't.png')
        self.fc.convert_file(str(self.dir / 't.png'), 'jpg')
        r, g, b = Image.open(self.dir / 't.jpg').getpixel((5, 5))
        self.assertGreater(min(r, g, b), 245)

    def test_ico_has_all_sizes_and_keeps_aspect(self):
        img = Image.new('RGBA', (300, 150), (0, 0, 255, 255))
        img.save(self.dir / 'wide.png')
        self.fc.convert_file(str(self.dir / 'wide.png'), 'ico')
        ico = Image.open(self.dir / 'wide.ico')
        self.assertIn((256, 256), ico.info['sizes'])
        self.assertIn((16, 16), ico.info['sizes'])
        ico.size = (256, 256)
        ico.load()
        self.assertEqual(ico.getpixel((128, 10))[3], 0)  # padding, not a stretched image
        self.assertEqual(ico.getpixel((128, 128))[3], 255)

    def test_svg_converts_without_cairo_and_keeps_aspect(self):
        (self.dir / 'a.svg').write_text(
            '<svg xmlns="http://www.w3.org/2000/svg" width="200" height="100">'
            '<rect width="200" height="100" fill="blue"/></svg>', encoding='utf-8')
        for mode in ('png', 'jpg', 'webp', 'ico'):
            ok, msg = self.fc.convert_file(str(self.dir / 'a.svg'), mode)
            self.assertTrue(ok, f'{mode}: {msg}')
        self.assertEqual(Image.open(self.dir / 'a.png').size, (1024, 512))

    def test_exif_orientation_is_applied(self):
        img = Image.new('RGB', (40, 20), 'red')
        exif = img.getexif()
        exif[0x0112] = 6  # rotate 90° CW on display
        img.save(self.dir / 'phone.jpg', exif=exif.tobytes())
        self.fc.convert_file(str(self.dir / 'phone.jpg'), 'png')
        self.assertEqual(Image.open(self.dir / 'phone.png').size, (20, 40))

    def test_heic_roundtrip(self):
        from pillow_heif import register_heif_opener
        register_heif_opener()
        Image.new('RGB', (32, 16), 'green').save(self.dir / 'p.heic')
        ok, msg = self.fc.convert_file(str(self.dir / 'p.heic'), 'jpg')
        self.assertTrue(ok, msg)
        self.assertEqual(Image.open(self.dir / 'p.jpg').size, (32, 16))

    def test_cmyk_jpeg_to_png(self):
        Image.new('CMYK', (8, 8), (0, 255, 255, 0)).save(self.dir / 'cmyk.jpg')
        ok, msg = self.fc.convert_file(str(self.dir / 'cmyk.jpg'), 'png')
        self.assertTrue(ok, msg)
        self.assertEqual(Image.open(self.dir / 'cmyk.png').mode, 'RGB')


class TestData(TempDirTestCase):
    @classmethod
    def setUpClass(cls):
        cls.fc = FileConverter()

    def test_bom_json(self):
        (self.dir / 'b.json').write_bytes('\ufeff[{"a": 1}]'.encode('utf-8'))
        ok, msg = self.fc.convert_file(str(self.dir / 'b.json'), 'csv')
        self.assertTrue(ok, msg)

    def test_excel_semicolon_csv(self):
        (self.dir / 'e.csv').write_bytes('\ufeffname;age\nBob;3\n'.encode('utf-8'))
        self.fc.convert_file(str(self.dir / 'e.csv'), 'json')
        self.assertEqual(json.loads((self.dir / 'e.json').read_text(encoding='utf-8')), [{'name': 'Bob', 'age': '3'}])

    def test_wrapped_list_and_nested_values_to_csv(self):
        (self.dir / 'n.json').write_text('{"users": [{"name": "a", "tags": ["x"]}, {"name": "b", "extra": 1}]}')
        self.fc.convert_file(str(self.dir / 'n.json'), 'csv')
        lines = (self.dir / 'n.csv').read_text(encoding='utf-8-sig').splitlines()
        self.assertEqual(lines[0], 'name,tags,extra')  # first-seen order, not alphabetical
        self.assertIn('"[""x""]"', lines[1])

    def test_yaml_dates_and_yml_extension(self):
        (self.dir / 'd.yml').write_text('when: 2024-01-02\n', encoding='utf-8')
        ok, msg = self.fc.convert_file(str(self.dir / 'd.yml'), 'json')
        self.assertTrue(ok, msg)
        self.assertEqual(json.loads((self.dir / 'd.json').read_text(encoding='utf-8')), {'when': '2024-01-02'})

    def test_json_to_xml_is_always_well_formed(self):
        (self.dir / 'k.json').write_text('{"1st": 1, "a b@c": null, "@id": 3, "#text": "t", "ok": [true, false]}')
        self.fc.convert_file(str(self.dir / 'k.json'), 'xml')
        root = ET.parse(self.dir / 'k.xml').getroot()
        self.assertEqual(root.get('id'), '3')
        self.assertEqual([e.text for e in root.findall('ok')], ['true', 'false'])

    def test_xml_json_xml_roundtrip(self):
        shutil.copy(helpers.SAMPLES / 'sample.xml', self.dir / 's.xml')
        self.fc.convert_file(str(self.dir / 's.xml'), 'json')
        self.fc.convert_file(str(self.dir / 's.json'), 'xml')
        items = ET.parse(self.dir / 's(1).xml').getroot().findall('item')
        self.assertEqual([i.get('id') for i in items], ['1', '2'])


class TestFolders(TempDirTestCase):
    def setUp(self):
        super().setUp()
        self.folder = self.dir / 'Holiday'
        self.folder.mkdir()
        for name, size in (('IMG_10.JPG', (400, 300)), ('IMG_2.jpg', (300, 400)), ('img_1.PNG', (50, 50))):
            Image.new('RGB', size, 'red').save(self.folder / name)
        (self.folder / 'notes.txt').write_text('ignore me')

    def test_folder_to_pdf_case_insensitive_natural_order(self):
        result = FileConverter().convert(str(self.folder), 'folderpdf')
        self.assertTrue(result.success, result.message)
        self.assertEqual(result.output, self.dir / 'Holiday.pdf')
        import pymupdf
        with pymupdf.open(str(result.output)) as doc:
            self.assertEqual(doc.page_count, 3)
            self.assertGreater(doc[2].rect.width, doc[2].rect.height)  # IMG_10 is landscape -> landscape page

    def test_folder_to_gif(self):
        result = FileConverter().convert(str(self.folder), 'foldergif')
        self.assertTrue(result.success, result.message)
        with Image.open(result.output) as gif:
            self.assertEqual(gif.n_frames, 3)


def _ffmpeg(*args):
    subprocess.run([get_ffmpeg_exe(), '-hide_banner', '-loglevel', 'error', '-y', *args], check=True)


def _video_info(path) -> str:
    return subprocess.run([get_ffmpeg_exe(), '-hide_banner', '-i', str(path)], capture_output=True, text=True).stderr


class TestMedia(TempDirTestCase):
    """Uses the FFmpeg bundled with imageio-ffmpeg — no system install needed."""

    @classmethod
    def setUpClass(cls):
        cls.fc = FileConverter()

    def setUp(self):
        super().setUp()
        self.video = self.dir / 'clip.mp4'
        _ffmpeg('-f', 'lavfi', '-i', 'testsrc=size=161x121:rate=10:duration=1',
                '-f', 'lavfi', '-i', 'sine=duration=1', '-shortest', '-pix_fmt', 'yuv444p', str(self.video))

    def test_video_to_audio_formats(self):
        for mode in ('mp3', 'wav', 'flac', 'aac', 'ogg'):
            ok, msg = self.fc.convert_file(str(self.video), mode)
            self.assertTrue(ok, f'{mode}: {msg}')
            self.assertGreater((self.dir / f'clip.{mode}').stat().st_size, 0)

    def test_video_without_audio_has_clear_error(self):
        silent = self.dir / 'silent.mp4'
        _ffmpeg('-f', 'lavfi', '-i', 'testsrc=size=64x64:rate=5:duration=1', '-pix_fmt', 'yuv420p', str(silent))
        ok, msg = self.fc.convert_file(str(silent), 'mp3')
        self.assertFalse(ok)
        self.assertIn('no audio', msg)
        self.assertFalse((self.dir / 'silent.mp3').exists())

    def test_video_to_gif_and_back_to_playable_mp4(self):
        ok, msg = self.fc.convert_file(str(self.video), 'gif')
        self.assertTrue(ok, msg)
        ok, msg = self.fc.convert_file(str(self.dir / 'clip.gif'), 'mp4')
        self.assertTrue(ok, msg)
        info = _video_info(self.dir / 'clip(1).mp4')
        self.assertIn('yuv420p', info)  # plays in Windows' built-in player
        self.assertIn('160x120', info)  # odd source size made even

    def test_mov_to_mp4(self):
        mov = self.dir / 'phone.mov'
        shutil.copy(self.video, mov)
        ok, msg = self.fc.convert_file(str(mov), 'mp4')
        self.assertTrue(ok, msg)

    def test_gif_frames_go_into_subfolder(self):
        frames = [Image.new('RGB', (8, 8), c) for c in ('red', 'green', 'blue')]
        frames[0].save(self.dir / 'a.gif', save_all=True, append_images=frames[1:], duration=100)
        result = self.fc.convert(str(self.dir / 'a.gif'), 'pngs')
        self.assertTrue(result.success, result.message)
        self.assertEqual(result.output, self.dir / 'a_frames')
        self.assertEqual(sorted(p.name for p in result.output.iterdir()), ['frame_0.png', 'frame_1.png', 'frame_2.png'])
        self.assertEqual(sorted(p.name for p in self.dir.iterdir() if p.suffix == '.png'), [])


class TestJobQueue(TempDirTestCase):
    def test_lock_is_exclusive(self):
        a, b = jobqueue.WorkerLock(self.dir / 'l'), jobqueue.WorkerLock(self.dir / 'l')
        self.assertTrue(a.acquire())
        self.assertFalse(b.acquire())
        a.release()
        self.assertTrue(b.acquire())
        b.release()

    def test_submit_take_roundtrip_and_stale_jobs_dropped(self):
        jobqueue.submit('C:/x.png', 'jpg')
        stale = jobqueue.QUEUE_DIR / 'old.job'
        stale.write_text(json.dumps({'path': 'C:/old.png', 'mode': 'jpg', 'time': time.time() - 10_000}))
        self.assertEqual(jobqueue.take_all(), [('C:/x.png', 'jpg')])
        self.assertFalse(jobqueue.has_jobs())

    def test_many_processes_one_batch(self):
        """Simulates Explorer starting one converter process per selected file."""
        files = []
        for i in range(6):
            p = self.dir / f'f{i}.png'
            Image.new('RGB', (20, 20), 'red').save(p)
            files.append(p)
        env = {**os.environ, 'PYTHONPATH': str(helpers.ROOT)}
        procs = [subprocess.Popen([sys.executable, str(helpers.ROOT / 'converter.py'), str(f), 'jpg'], env=env,
                                  stdout=subprocess.PIPE, stderr=subprocess.STDOUT) for f in files]
        outputs = [p.communicate(timeout=120)[0].decode(errors='ignore') for p in procs]
        self.assertTrue(all(p.returncode == 0 for p in procs), outputs)
        for f in files:
            self.assertTrue(f.with_suffix('.jpg').exists(), f)
        self.assertFalse(jobqueue.has_jobs())

    def test_main_without_args_does_not_crash(self):
        import converter
        shown = []
        original = converter.show_about
        converter.show_about = lambda: shown.append(True)
        try:
            self.assertEqual(converter_main(['converter.exe']), 0)
        finally:
            converter.show_about = original
        self.assertEqual(shown, [True])


class TestFormatsConsistency(unittest.TestCase):
    def test_every_menu_entry_has_a_converter(self):
        fc = FileConverter()
        for src, targets in CONVERSION_MAP.items():
            self.assertIn(src, fc._routers, f'.{src} in menu but no converter claims it')
            self.assertNotIn(src, targets, f'.{src} -> .{src} offered')
        self.assertEqual(FOLDER_TARGETS, ['pdf', 'gif'])

    def test_registry_fragment_covers_map(self):
        sys.path.insert(0, str(helpers.ROOT / 'installer'))
        import generate_registry_iss
        text = generate_registry_iss.render()
        for src, targets in CONVERSION_MAP.items():
            self.assertIn(f'SystemFileAssociations\\.{src}\\shell\\PocketConverter', text)
        self.assertIn('Tasks: contextmenu', text)


if __name__ == '__main__':
    unittest.main(verbosity=2)
